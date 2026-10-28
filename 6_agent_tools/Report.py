# /home/yzt/AgentDNS/src/docstore/Report.py
# -*- coding: utf-8 -*-

import os
import re
import json
import asyncio
from typing import Any, Optional, List, Union

from dotenv import load_dotenv
load_dotenv("/home/yzt/AgentDNS/src/react/.env", override=True)

from langchain_openai import ChatOpenAI
from langchain_core.messages import HumanMessage, SystemMessage

# 可选：reward_manager
try:
    import src.utils.reward_manager as reward_manager  # type: ignore
except Exception:
    reward_manager = None


# -------------------------
# Utils
# -------------------------
def _to_text(x: Any) -> str:
    if x is None:
        return ""
    if isinstance(x, str):
        return x
    try:
        return json.dumps(x, ensure_ascii=False, default=str)
    except Exception:
        return str(x)


def _clean_str(s: str) -> str:
    """
    清洗 LLM 传参里的脏字符，典型问题：/path/test.mp4") 之类
    """
    s = _to_text(s).strip()
    # 去掉外层引号
    s = s.strip().strip('"').strip("'").strip()
    # 去掉末尾多余引号/括号/分号
    s = re.sub(r'[\s\)\];]+$', "", s)
    s = s.strip().strip('"').strip("'").strip()
    return s


def _extract_video_path(s: str) -> Optional[str]:
    s = _to_text(s)
    m = re.search(r"VIDEO_PATH\s*=\s*([^\s`]+)", s)
    if m:
        return _clean_str(m.group(1))

    m2 = re.search(r"(/[^ \n\t`]+?\.(?:mp4|mov|avi|mkv|webm))", s, flags=re.I)
    if m2:
        return _clean_str(m2.group(1))
    return None


def _extract_prev_step_summary(text: str) -> Optional[str]:
    """
    从传入文本里提取 PREV_STEP_SUMMARY 段落（workflow 最关键的上一步摘要）
    """
    t = _to_text(text)
    if "PREV_STEP_SUMMARY" not in t:
        return None
    m = re.search(r"PREV_STEP_SUMMARY\s*:\s*(.*)$", t, flags=re.S)
    if not m:
        return None
    block = m.group(1).strip()
    return block if block else None


def _read_json_file(path: str) -> Any:
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def _read_text_file(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _normalize_records_to_text(records: Any, limit: int = 120) -> str:
    """
    把各种 logs 结构统一成“系统记录文本”
    - 支持 list[dict] / dict / str
    """
    if isinstance(records, str):
        return records

    chunks: List[str] = []
    if isinstance(records, dict):
        for k in ["Agent scratchpad", "Question", "Raw Answer", "Observation", "Result", "answer", "text", "summary"]:
            if k in records and records[k]:
                chunks.append(f"{k}:\n{_to_text(records[k])}")
        if not chunks:
            chunks.append(_to_text(records))
        return "\n\n".join(chunks)

    if isinstance(records, list):
        for i, item in enumerate(records[:limit]):
            if isinstance(item, dict):
                sub = []
                for k in ["Agent scratchpad", "Question", "Raw Answer", "agent_name", "tool", "observation", "result", "time", "summary"]:
                    if k in item and item[k]:
                        sub.append(f"{k}: {_to_text(item[k])}")
                if not sub:
                    sub.append(_to_text(item))
                chunks.append(f"[Record {i+1}]\n" + "\n".join(sub))
            else:
                chunks.append(f"[Record {i+1}]\n{_to_text(item)}")
        return "\n\n".join(chunks)

    return _to_text(records)


# -------------------------
# Agent
# -------------------------
class ReportAgent:
    """
    ✅ 类名必须叫 ReportAgent
    ✅ 必须包含：read_file_async / generate_report_async / get_image_paths_async / write_report_async
    """

    def __init__(self, model: str | None = None):
        model = model or os.getenv("MODEL_NAME", "qwen-plus")

        api_key = os.getenv("DASHSCOPE_API_KEY")
        if not api_key:
            raise ValueError("Missing env var: DASHSCOPE_API_KEY")

        base_url = os.getenv("DASHSCOPE_BASE_URL", "https://dashscope.aliyuncs.com/compatible-mode/v1")

        self.chat = ChatOpenAI(
            model=model,
            temperature=0,
            max_tokens=4096,  # 报告不需要太长，避免慢/超长
            api_key=api_key,
            base_url=base_url,
        )

        self.all_tasks_log_path = os.getenv(
            "ALL_TASKS_LOG_PATH",
            "/home/yzt/LLMCompiler/result/logs/all_tasks_logs.json"
        )

    # -------------------------
    # (A) ReadFile：支持 “$PREV_STEP_SUMMARY / 完整Question(含PREV) / 视频路径 / 普通文件”
    # -------------------------
    def read_file_sync(self, file_path_or_text: str) -> str:
        try:
            raw = _to_text(file_path_or_text)
            if not raw or not raw.strip():
                return "发生错误: empty_input"

            s_clean = _clean_str(raw)

            # 0) 允许直接喂 env 占位符
            if s_clean == "$PREV_STEP_SUMMARY":
                prev_env = (os.getenv("PREV_STEP_SUMMARY") or "").strip()
                return prev_env if prev_env else "发生错误: PREV_STEP_SUMMARY 为空（未由 workflow runner 注入）"

            # 1) 如果传进来的是完整 Question（含 PREV_STEP_SUMMARY），优先直接返回摘要（最稳，不串台）
            prev_in_text = _extract_prev_step_summary(raw)
            if prev_in_text:
                return "=== PREV_STEP_SUMMARY (preferred) ===\n" + prev_in_text

            # 2) VIDEO_PATH 或直接 video 文件路径
            video_path = _extract_video_path(raw)
            if not video_path and re.search(r"\.(mp4|mov|avi|mkv|webm)$", s_clean, re.I):
                video_path = s_clean

            if video_path:
                # 优先吃 workflow 注入的 env（你们 workflow 跑出来的上一步摘要）
                prev_env = (os.getenv("PREV_STEP_SUMMARY") or "").strip()
                if prev_env:
                    return f"VIDEO_PATH={video_path}\n\nPREV_STEP_SUMMARY:\n{prev_env}"
                # 没有 env 才尝试从 all_tasks_logs 捞
                return self._read_logs_for_video(video_path)

            # 3) 普通文件
            if os.path.isfile(s_clean):
                if s_clean.lower().endswith(".json"):
                    data = _read_json_file(s_clean)
                    return _normalize_records_to_text(data)
                return _read_text_file(s_clean)

            # 4) 当普通文本
            return raw.strip()

        except Exception as e:
            return f"发生错误: {e}"

    async def read_file_async(self, file_path_or_text: str) -> str:
        loop = asyncio.get_event_loop()
        return await loop.run_in_executor(None, self.read_file_sync, file_path_or_text)

    # -------------------------
    # (B) 为某个视频找“系统记录”（仅在你真的把 video_path 写入 all_tasks_logs 时才有用）
    # -------------------------
    def _read_logs_for_video(self, video_path: str) -> str:
        vp = str(video_path).strip()

        # 1) 兼容旧逻辑：video.mp4.json
        legacy_json = vp + ".json"
        if os.path.isfile(legacy_json):
            try:
                data = _read_json_file(legacy_json)
                return _normalize_records_to_text(data)
            except Exception:
                pass

        # 2) 从 all_tasks_logs 过滤
        logp = self.all_tasks_log_path
        candidates = [logp]
        if not os.path.isfile(logp) and os.path.isfile(logp + ".json"):
            candidates.insert(0, logp + ".json")

        chosen = None
        for p in candidates:
            if os.path.isfile(p):
                chosen = p
                break

        if not chosen:
            return (
                "发生错误: 找不到全局任务日志 ALL_TASKS_LOG_PATH。\n"
                f"当前 video_path={vp}\n"
                f"请设置 env: ALL_TASKS_LOG_PATH 指向 all_tasks_logs.json 或文本日志文件。"
            )

        fn = os.path.basename(vp)

        if chosen.lower().endswith(".json"):
            try:
                data = _read_json_file(chosen)
            except Exception as e:
                return f"发生错误: 读取 {chosen} 失败: {e}"

            text_blob = _to_text(data)
            picked: List[Any] = []

            if isinstance(data, list):
                for item in data:
                    blob = _to_text(item)
                    if vp in blob or fn in blob:
                        picked.append(item)
            else:
                if vp in text_blob or fn in text_blob:
                    picked = [data]

            if not picked:
                return (
                    f"VIDEO_PATH={vp}\n\n"
                    "=== Filtered System Records ===\n"
                    "未在 ALL_TASKS_LOG_PATH 中找到与该 VIDEO_PATH 匹配的记录。\n"
                    "建议：\n"
                    "1) workflow runner 写 all_tasks_logs 时加入 video_path 字段；或\n"
                    "2) 给 ReportAgent 注入 PREV_STEP_SUMMARY（推荐）。\n"
                )

            return (
                f"VIDEO_PATH={vp}\n\n"
                "=== Filtered System Records ===\n"
                + _normalize_records_to_text(picked, limit=80)
            )

        # 纯文本形式
        try:
            txt = _read_text_file(chosen)
        except Exception as e:
            return f"发生错误: 读取 {chosen} 失败: {e}"

        lines = txt.splitlines()
        hit_lines = [ln for ln in lines if (vp in ln or fn in ln)]
        if not hit_lines:
            hit_lines = lines[-200:]

        return (
            f"VIDEO_PATH={vp}\n\n"
            "=== Filtered System Records (text) ===\n"
            + "\n".join(hit_lines[-400:])
        )

    # -------------------------
    # (C) 生成报告（中文）
    # -------------------------
    def generate_report_sync(self, logs: str) -> str:
        logs = _to_text(logs).strip()

        messages = [
            SystemMessage(
                content=(
                    "你是事故分析报告撰写专家，请用【中文】输出结构化报告。\n"
                    "硬性规则：\n"
                    "1) 只能使用输入中的 system records / PREV_STEP_SUMMARY，禁止编造。\n"
                    "2) 没提到就写“未知/未提及”，不要猜。\n"
                    "3) 允许用要点列表，信息尽量密。\n"
                    "输出必须严格包含以下标题（且仅这些）：\n"
                    "标题：\n"
                    "一、事故概述：\n"
                    "二、关键证据（来自视频/关键帧）：\n"
                    "三、人员情况：\n"
                    "四、处置与建议：\n"
                    "五、信息缺口与下一步：\n"
                )
            ),
            HumanMessage(content=f"System records / PREV_STEP_SUMMARY:\n{logs}\n\n现在生成报告。"),
        ]

        resp = self.chat.invoke(messages)
        out = _to_text(resp.content).strip()

        if not re.search(r"^\s*标题[:：]", out, flags=re.M):
            out = "标题：事故分析报告\n" + out
        return out

    async def generate_report_async(self, logs: str) -> str:
        loop = asyncio.get_event_loop()
        return await loop.run_in_executor(None, self.generate_report_sync, logs)

    # -------------------------
    # (D) 提取图片路径
    # -------------------------
    def get_image_paths_sync(self, text: str) -> List[str]:
        # 0) 优先：workflow runner 注入的上一阶段图片列表（最稳）
        prev_imgs_env = (os.getenv("PREV_STEP_IMAGES") or "").strip()
        if prev_imgs_env:
            try:
                arr = json.loads(prev_imgs_env)
                if isinstance(arr, list):
                    arr = [_clean_str(str(x)) for x in arr]
                    arr = [p for p in arr if p and os.path.isfile(p)]
                    if arr:
                        print("已获取图片路径(ENV PREV_STEP_IMAGES)：[" + ",".join(arr) + "]")
                        return arr
            except Exception:
                pass

        raw = _to_text(text)

        # 1) 次优：从 PREV_STEP_SUMMARY（env 或传入文本）里抽 jpg/png 路径
        prev_env = (os.getenv("PREV_STEP_SUMMARY") or "").strip()
        prev_in_text = _extract_prev_step_summary(raw) or ""
        merged = "\n".join([prev_env, prev_in_text, raw])

        # 2) 从 merged 文本里提取图片路径，并清洗脏字符
        paths = re.findall(r"(\/[^\s]+?\.(?:jpg|jpeg|png|webp))", merged, flags=re.I)
        out = []
        for p in paths:
            p2 = _clean_str(p)
            if p2 and os.path.isfile(p2):
                out.append(p2)

        # 去重
        out = list(dict.fromkeys(out))

        print("已获取图片路径：[" + ",".join(out) + "]")
        return out

    async def get_image_paths_async(self, text: str) -> List[str]:
        loop = asyncio.get_event_loop()
        return await loop.run_in_executor(None, self.get_image_paths_sync, text)

    # -------------------------
    # (E) 写入 Word（支持中文标题）
    # -------------------------
    def write_report_sync(self, report_content: str, image_paths: Optional[Union[str, List[str]]] = None) -> str:
        try:
            from docx import Document
            from docx.shared import Inches, Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_PARAGRAPH_ALIGNMENT
        except Exception as e:
            return f"[WARN] python-docx not available: {e}"

        report_content = _to_text(report_content).replace("*", "").strip()
        doc = Document()

        title = "事故分析报告"
        m_cn = re.search(r"^\s*标题[:：]\s*(.+)\s*$", report_content, flags=re.M)
        m_en = re.search(r"^\s*Title:\s*(.+)\s*$", report_content, flags=re.M)
        if m_cn:
            title = m_cn.group(1).strip()[:120]
        elif m_en:
            title = m_en.group(1).strip()[:120]

        h = doc.add_heading(title, level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for para in report_content.split("\n"):
            para = para.strip()
            if not para:
                continue
            p = doc.add_paragraph(para)
            fmt = p.paragraph_format
            fmt.line_spacing = Pt(14)
            fmt.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
            fmt.space_after = Pt(6)

        # 图片
        paths: List[str] = []
        if image_paths:
            if isinstance(image_paths, str):
                s = image_paths.strip()
                try:
                    parsed = json.loads(s)
                    if isinstance(parsed, list):
                        paths = [str(x).strip() for x in parsed]
                    else:
                        paths = [x.strip() for x in s.strip("[]").replace("'", "").split(",") if x.strip()]
                except Exception:
                    paths = [x.strip() for x in s.strip("[]").replace("'", "").split(",") if x.strip()]
            elif isinstance(image_paths, list):
                paths = [str(x).strip() for x in image_paths if str(x).strip()]

        if paths:
            doc.add_heading("Images", level=2)
            for pth in paths[:12]:
                try:
                    if os.path.isfile(pth):
                        doc.add_picture(pth, width=Inches(3))
                        doc.add_paragraph(f"Image: {os.path.basename(pth)}")
                        if reward_manager:
                            try:
                                reward_manager.add_reward(5)
                            except Exception:
                                pass
                except Exception as e:
                    doc.add_paragraph(f"[WARN] cannot add image {pth}: {e}")

        out_dir = os.getenv("REPORT_OUTPUT_DIR", "/home/yzt/LLMCompiler/result/reports")
        os.makedirs(out_dir, exist_ok=True)

        safe = re.sub(r"[\\/:*?\"<>|]+", "_", title).strip() or "事故分析报告"
        file_path = os.path.join(out_dir, f"{safe}.docx")
        doc.save(file_path)
        return file_path

    async def write_report_async(self, report_content: str, image_paths: Optional[Union[str, List[str]]] = None) -> str:
        loop = asyncio.get_event_loop()
        return await loop.run_in_executor(None, self.write_report_sync, report_content, image_paths)


__all__ = ["ReportAgent"]
